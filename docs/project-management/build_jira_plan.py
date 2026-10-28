# -*- coding: utf-8 -*-
"""Đọc sheet "3. Master Plan" của nhóm (Template-Technical1.xlsx) và sinh ra:
  - jira_plan.json : cây Epic -> Task -> Subtask (summary, mô tả, người phụ trách, ngày bắt đầu/kết thúc, nhãn)
  - jira_import.csv: file dự phòng để nhập bằng chức năng CSV import của Jira (Issue ID / Parent ID)
  - KE_HOACH_JIRA.md: bảng đối chiếu để rà soát trước khi tạo issue
Chạy: python -I build_jira_plan.py "<đường dẫn Template-Technical1.xlsx>" "<thư mục xuất>"
"""
import sys, re, json, csv, datetime, collections
import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
SRC = sys.argv[1]
OUT_DIR = sys.argv[2] if len(sys.argv) > 2 else "."

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["3. Master Plan"]

MEMBERS = ["Lê Phú Toàn", "Lê Phương Linh", "Trần Đăng Việt", "Phạm Phương Anh"]
ROMAN = re.compile(r"^(I|II|III|IV|V|VI|VII|VIII|IX|X)$")
TASK = re.compile(r"^WMS-(\d{2})$")
SUB = re.compile(r"^WMS-(\d{2})\.(\d+)$")
TYPE = {"epic": "Epic", "task": "Task", "subtask": "Subtask"}


def clean(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s).replace("\n", " ")).strip()


def to_iso(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    return None


def fmt(iso):
    y, m, dd = iso.split("-")
    return f"{dd}/{m}/{y}"


def sprint_labels(text):
    """Trích nhãn giai đoạn từ phần [..] đầu cột Tóm tắt, ví dụ [Chuẩn bị & Sprint 1–3]."""
    m = re.match(r"\s*\[([^\]]+)\]", text or "")
    tag = m.group(1) if m else ""
    labels = []
    if "Chuẩn bị" in tag:
        labels.append("Chuan-bi")
    for a, b in re.findall(r"Sprint\s*(\d)(?:\s*[–-]\s*(\d))?", tag):
        lo, hi = int(a), int(b) if b else int(a)
        for n in range(lo, hi + 1):
            lab = f"Sprint-{n}"
            if lab not in labels:
                labels.append(lab)
    if "UAT" in tag or "Deploy" in tag:
        labels.append("UAT-Deploy")
    return labels, tag


def role_label(title):
    m = re.match(r"\s*\+?\s*\[([A-Za-z]+)\]", title or "")
    return m.group(1) if m else None


issues = []
epic = None
task = None
for r in range(9, ws.max_row + 1):
    stt = clean(ws[f"A{r}"].value)
    if not stt:
        continue
    title = clean(ws[f"B{r}"].value)
    summ = clean(ws[f"C{r}"].value)
    pic = clean(ws[f"D{r}"].value)
    collab = clean(ws[f"E{r}"].value)
    prog = ws[f"F{r}"].value or 0
    start, end = to_iso(ws[f"G{r}"].value), to_iso(ws[f"H{r}"].value)
    days = ws[f"I{r}"].value
    if not isinstance(days, (int, float)) and start and end:
        days = (datetime.date.fromisoformat(end) - datetime.date.fromisoformat(start)).days + 1
    base = dict(row=r, wms=stt, title_raw=title, summary_raw=summ, assignee=pic, collab=collab,
                progress=float(prog) if isinstance(prog, (int, float)) else 0.0,
                start=start, due=end, days=int(days) if isinstance(days, (int, float)) else None)
    if ROMAN.match(stt):
        labels, tag = sprint_labels(summ)
        epic = dict(base, level="epic", parent=None, summary=title, labels=labels, phase=tag)
        task = None
        issues.append(epic)
    elif TASK.match(stt):
        labels, tag = sprint_labels(summ)
        name = re.sub(r"^Task\s+WMS-\d{2}:\s*", "", title)
        task = dict(base, level="task", parent=epic["wms"], summary=f"{stt}: {name}", labels=labels, phase=tag)
        issues.append(task)
    elif SUB.match(stt):
        name = re.sub(r"^\+\s*", "", title)
        role = role_label(name)
        labels = list(task["labels"]) + ([role] if role else [])
        issues.append(dict(base, level="subtask", parent=task["wms"], summary=f"{stt} {name}",
                           labels=labels, phase=task["phase"]))
    else:
        print("Bỏ qua dòng không nhận dạng:", r, stt)

# Kiểm tra tính nhất quán
cnt = collections.Counter(i["level"] for i in issues)
assert cnt == {"epic": 9, "task": 21, "subtask": 86}, cnt
by_wms = {i["wms"]: i for i in issues}
for i in issues:
    assert i["start"] and i["due"], (i["wms"], "thiếu ngày")
    assert i["start"] <= i["due"], (i["wms"], "ngày bắt đầu sau ngày kết thúc")
    assert len(i["summary"]) <= 255, (i["wms"], "summary quá dài")
    assert i["assignee"] in MEMBERS or i["assignee"] == "Cả 4 thành viên", (i["wms"], i["assignee"])
    if i["parent"]:
        p = by_wms[i["parent"]]
        if not (p["start"] <= i["start"] and i["due"] <= p["due"]):
            print(f"Cảnh báo: {i['wms']} ({i['start']}..{i['due']}) ngoài khoảng của {p['wms']} ({p['start']}..{p['due']})")

for i in issues:
    lines = [("Mã Master Plan", i["wms"]),
             ("Nội dung", i["summary_raw"] or i["title_raw"]),
             ("Phụ trách", i["assignee"]),
             ("Phối hợp", i["collab"] or "-"),
             ("Kế hoạch", f"{fmt(i['start'])} - {fmt(i['due'])} ({i['days']} ngày)"),
             ("Giai đoạn", i["phase"] or "-")]
    if i["parent"]:
        lines.append(("Thuộc", i["parent"]))
    i["description_lines"] = lines

plan = dict(source=SRC, generated=datetime.datetime.now().isoformat(timespec="seconds"),
            project_key="TASK", members=MEMBERS, counts=dict(cnt),
            period=dict(start=min(i["start"] for i in issues), end=max(i["due"] for i in issues)),
            issues=issues)
with open(f"{OUT_DIR}/jira_plan.json", "w", encoding="utf-8") as f:
    json.dump(plan, f, ensure_ascii=False, indent=1)

# CSV dự phòng cho Jira CSV import (ngày dạng dd/MM/yyyy, khai báo định dạng khi import)
ids = {i["wms"]: n + 1 for n, i in enumerate(issues)}
maxlab = max(len(i["labels"]) for i in issues)
with open(f"{OUT_DIR}/jira_import.csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(["Issue ID", "Parent ID", "Issue Type", "Summary", "Description", "Assignee",
                "Start Date", "Due Date"] + ["Labels"] * maxlab)
    for i in issues:
        desc = "\n".join(f"*{k}:* {v}" for k, v in i["description_lines"])
        labs = i["labels"] + [""] * (maxlab - len(i["labels"]))
        w.writerow([ids[i["wms"]], ids.get(i["parent"], "") if i["parent"] else "", TYPE[i["level"]],
                    i["summary"], desc, i["assignee"], fmt(i["start"]), fmt(i["due"])] + labs)

# Bảng rà soát
per = collections.Counter(i["assignee"] for i in issues if i["level"] == "subtask")
with open(f"{OUT_DIR}/KE_HOACH_JIRA.md", "w", encoding="utf-8") as f:
    f.write("# Kế hoạch đưa Master Plan lên Jira (dự án TASK)\n\n")
    f.write(f"Nguồn: Template-Technical1.xlsx, sheet 3. Master Plan (người lập: Phạm Phương Anh). "
            f"Giai đoạn {fmt(plan['period']['start'])} - {fmt(plan['period']['end'])}.\n\n")
    f.write("Số lượng: 9 Epic, 21 Task, 86 Subtask (tổng 116 issue).\n\n")
    f.write("## Số subtask theo người phụ trách\n\n| Thành viên | Subtask |\n|---|---|\n")
    for k, v in per.most_common():
        f.write(f"| {k} | {v} |\n")
    f.write("\n## Cây công việc\n\n| Mã | Loại | Summary trên Jira | Phụ trách | Bắt đầu | Kết thúc | Nhãn |\n|---|---|---|---|---|---|---|\n")
    for i in issues:
        ind = {"epic": "", "task": "- ", "subtask": "- - "}[i["level"]]
        f.write(f"| {i['wms']} | {TYPE[i['level']]} | {ind}{i['summary']} | {i['assignee']} | {fmt(i['start'])} | {fmt(i['due'])} | {', '.join(i['labels'])} |\n")

print("OK", dict(cnt), "giai đoạn", plan["period"])
print("Subtask theo người:", dict(per))
print("Nhãn:", sorted({l for i in issues for l in i["labels"]}))
