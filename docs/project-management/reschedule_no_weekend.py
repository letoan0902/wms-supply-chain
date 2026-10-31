# -*- coding: utf-8 -*-
"""Sửa Master Plan để không giao việc vào thứ Bảy, Chủ nhật, giữ nguyên ngày kết thúc 23/11/2026.

Quy tắc (nhóm chốt 31/10/2026):
- Giữ 4 đợt của kế hoạch: 27/10-02/11, 03/11-09/11, 10/11-16/11, 17/11-23/11. Mỗi đợt 7 ngày lịch quy về 5 ngày làm
  việc (thứ Hai đến thứ Sáu), công việc giữ nguyên thứ tự và vị trí tương đối trong đợt.
- Đợt 1 đã trôi qua gần hết: ngày đến 30/10 giữ nguyên; việc rơi vào 31/10, 01/11 chuyển sang 02/11. Riêng WMS-04 và
  WMS-08 (vốn được ghi là Sprint 1 và Sprint 2) trải 4 subtask lần lượt vào 02, 03, 04, 05/11 để không dồn 4 việc vào
  một ngày.
- Task và Epic lấy ngày bắt đầu sớm nhất và kết thúc muộn nhất của các việc con.
Đầu ra: xlsx mới (file gốc giữ nguyên), csv đối chiếu cũ - mới, cập nhật jira_plan.json (start, due, days, dòng Kế hoạch).
Chạy: python -I reschedule_no_weekend.py <xlsx gốc> <xlsx mới> <csv đối chiếu> <jira_plan.json>
"""
import sys, re, csv, json, datetime as dt
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
SRC, OUT, CSV, PLAN = sys.argv[1:5]
D = dt.date
WINDOWS = [(D(2026, 10, 27), D(2026, 11, 2)), (D(2026, 11, 3), D(2026, 11, 9)),
           (D(2026, 11, 10), D(2026, 11, 16)), (D(2026, 11, 17), D(2026, 11, 23))]
IDX = [0, 1, 1, 2, 3, 3, 4]
THU = ["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ nhật"]
SPREAD = {"WMS-04": [D(2026, 11, 2), D(2026, 11, 3), D(2026, 11, 4), D(2026, 11, 5)],
          "WMS-08": [D(2026, 11, 2), D(2026, 11, 3), D(2026, 11, 4), D(2026, 11, 5)]}


def working_days(a, b):
    return [a + dt.timedelta(n) for n in range((b - a).days + 1) if (a + dt.timedelta(n)).weekday() < 5]


def remap(d):
    w0, w1 = WINDOWS[0]
    if w0 <= d <= w1:  # đợt 1: quá khứ giữ nguyên, cuối tuần dồn về thứ Hai 02/11
        return d if d <= D(2026, 10, 30) else D(2026, 11, 2)
    for w0, w1 in WINDOWS[1:]:
        if w0 <= d <= w1:
            return working_days(w0, w1)[IDX[(d - w0).days]]
    raise ValueError(f"Ngày {d} ngoài kế hoạch")


def fmt(d):
    return d.strftime("%d/%m/%Y")


wb = openpyxl.load_workbook(SRC)
wbv = openpyxl.load_workbook(SRC, data_only=True)
ws, wsv = wb["3. Master Plan"], wbv["3. Master Plan"]
ROMAN = re.compile(r"^(I|II|III|IV|V|VI|VII|VIII|IX)$")
items = []  # mọi dòng kế hoạch (epic, task, subtask) với ngày cũ
for r in range(9, ws.max_row + 1):
    stt = str(ws[f"A{r}"].value or "").strip()
    if not stt:
        continue
    g, h = wsv[f"G{r}"].value, wsv[f"H{r}"].value
    if not isinstance(g, dt.datetime) or not isinstance(h, dt.datetime):
        continue
    level = "epic" if ROMAN.match(stt) else ("task" if re.match(r"^WMS-\d\d$", stt) else "subtask")
    items.append(dict(row=r, stt=stt, level=level, title=str(ws[f"B{r}"].value or "").strip(),
                      pic=str(ws[f"D{r}"].value or "").strip(), s0=g.date(), e0=h.date(), d0=wsv[f"I{r}"].value,
                      formula_dates=not isinstance(ws[f"G{r}"].value, dt.datetime)))

# 1. Tính ngày mới cho subtask
cur_task, order_in_task = None, {}
for it in items:
    if it["level"] == "task":
        cur_task = it["stt"]
    if it["level"] == "subtask":
        order_in_task.setdefault(cur_task, []).append(it["stt"])
        it["parent"] = cur_task
for it in items:
    if it["level"] != "subtask":
        continue
    if it["parent"] in SPREAD:
        d = SPREAD[it["parent"]][order_in_task[it["parent"]].index(it["stt"])]
        it["s1"], it["e1"] = d, d
    else:
        it["s1"], it["e1"] = remap(it["s0"]), remap(it["e0"])
    assert it["s1"] <= it["e1"] and it["s1"].weekday() < 5 and it["e1"].weekday() < 5, it["stt"]
    it["d1"] = len(working_days(it["s1"], it["e1"]))

# 2. Task và Epic theo việc con
by_stt = {it["stt"]: it for it in items}
cur_epic = None
for it in items:
    if it["level"] == "epic":
        cur_epic = it["stt"]
    if it["level"] == "task":
        it["epic"] = cur_epic
for level in ("task", "epic"):
    for it in items:
        if it["level"] != level:
            continue
        kids = [k for k in items if (k["level"] == "subtask" and k.get("parent") == it["stt"]) or
                (k["level"] == "task" and k.get("epic") == it["stt"])]
        it["s1"], it["e1"] = min(k["s1"] for k in kids), max(k["e1"] for k in kids)
        it["d1"] = len(working_days(it["s1"], it["e1"]))

# 3. Ghi xlsx mới: ô ngày là giá trị thì ghi giá trị; ô công thức (MIN/MAX) giữ nguyên để Excel tự tính
for it in items:
    r = it["row"]
    if not it["formula_dates"]:
        ws[f"G{r}"].value = dt.datetime(it["s1"].year, it["s1"].month, it["s1"].day)
        ws[f"H{r}"].value = dt.datetime(it["e1"].year, it["e1"].month, it["e1"].day)
    if not (isinstance(ws[f"I{r}"].value, str) and str(ws[f"I{r}"].value).startswith("=")):
        ws[f"I{r}"].value = it["d1"]
ws["C8"].value = (str(ws["C8"].value or "") + " Lịch chỉ dùng ngày làm việc thứ Hai đến thứ Sáu, không giao việc vào "
                  "thứ Bảy và Chủ nhật, giữ ngày kết thúc 23/11/2026 (cập nhật 31/10/2026).").strip()
if not (isinstance(ws["I8"].value, str) and str(ws["I8"].value).startswith("=")):
    ws["I8"].value = len(working_days(D(2026, 10, 27), D(2026, 11, 23)))
ns = wb["4. Nhân sự dự án"]
for row in ns.iter_rows(min_row=1, max_row=ns.max_row):
    for c in row:
        if isinstance(c.value, str) and "28 ngày" in c.value:
            c.value = c.value.replace("28 ngày", "20 ngày làm việc")
wb.save(OUT)

# 4. CSV đối chiếu
with open(CSV, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(["Mã", "Loại", "Công việc", "Phụ trách", "Bắt đầu cũ", "Kết thúc cũ", "Số ngày cũ",
                "Bắt đầu mới", "Kết thúc mới", "Số ngày làm việc mới", "Thay đổi"])
    for it in items:
        w.writerow([it["stt"], it["level"], it["title"], it["pic"], fmt(it["s0"]), fmt(it["e0"]), it["d0"],
                    fmt(it["s1"]), fmt(it["e1"]), it["d1"], "x" if (it["s1"], it["e1"]) != (it["s0"], it["e0"]) else ""])

# 5. Cập nhật jira_plan.json
plan = json.load(open(PLAN, encoding="utf-8"))
for i in plan["issues"]:
    it = by_stt[i["wms"]]
    i["start"], i["due"], i["days"] = it["s1"].isoformat(), it["e1"].isoformat(), it["d1"]
    for line in i["description_lines"]:
        if line[0] == "Kế hoạch":
            line[1] = f"{fmt(it['s1'])} - {fmt(it['e1'])} ({it['d1']} ngày làm việc)" + (
                f"; Giai đoạn: {i['phase']}" if False else "")
plan["period"] = dict(start="2026-10-27", end="2026-11-23")
plan["note"] = "Lịch không có thứ Bảy, Chủ nhật, cập nhật 31/10/2026"
json.dump(plan, open(PLAN, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

# 6. Tóm tắt
subs = [it for it in items if it["level"] == "subtask"]
wk = [it for it in subs if it["s0"].weekday() >= 5 or it["e0"].weekday() >= 5]
print(f"{len(items)} dòng, {len(subs)} subtask; subtask có ngày cuối tuần: {len(wk)}; đổi ngày: "
      f"{sum(1 for it in items if (it['s1'], it['e1']) != (it['s0'], it['e0']))}/{len(items)}")
print("Ngày công subtask: cũ", sum(it["d0"] for it in subs if isinstance(it["d0"], (int, float))), "-> mới", sum(it["d1"] for it in subs))
print("Còn cuối tuần:", sum(1 for it in items if it["s1"].weekday() >= 5 or it["e1"].weekday() >= 5))
print("Tải theo người và ngày (số subtask có ngày đó trong khoảng):")
import collections
load = collections.defaultdict(lambda: collections.Counter())
for it in subs:
    for d in working_days(it["s1"], it["e1"]):
        load[it["pic"]][d] += 1
for p, c in load.items():
    heavy = sorted((d, n) for d, n in c.items() if n >= 3)
    print(f"  {p:16} ngày có >=3 việc: " + (", ".join(f"{fmt(d)}:{n}" for d, n in heavy) if heavy else "không"))
print("Epic / task mới:")
for it in items:
    if it["level"] != "subtask":
        print(f"  {it['stt']:7} {fmt(it['s0'])}-{fmt(it['e0'])} -> {fmt(it['s1'])}-{fmt(it['e1'])} ({it['d1']} ngày lv)")
